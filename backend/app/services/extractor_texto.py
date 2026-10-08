"""
Extrae el texto plano de una hoja de vida subida como PDF, DOCX o TXT.
"""

import io
import re

TAMANO_MAXIMO = 5 * 1024 * 1024  # 5 MB
EXTENSIONES = (".pdf", ".docx", ".txt")


class FormatoNoSoportado(ValueError):
    pass


class ArchivoIlegible(ValueError):
    pass


def _limpiar(texto: str) -> str:
    texto = texto.replace("\x00", "")
    texto = re.sub(r"[ \t]+", " ", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()


def _pdf(contenido: bytes) -> str:
    from pypdf import PdfReader

    lector = PdfReader(io.BytesIO(contenido))
    return "\n".join(pagina.extract_text() or "" for pagina in lector.pages)


def _docx(contenido: bytes) -> str:
    from docx import Document

    documento = Document(io.BytesIO(contenido))
    partes = [p.text for p in documento.paragraphs]
    # Muchas hojas de vida usan tablas para organizar secciones
    for tabla in documento.tables:
        for fila in tabla.rows:
            partes.append(" | ".join(celda.text for celda in fila.cells))
    return "\n".join(partes)


def _txt(contenido: bytes) -> str:
    for codificacion in ("utf-8", "latin-1"):
        try:
            return contenido.decode(codificacion)
        except UnicodeDecodeError:
            continue
    raise ArchivoIlegible("No se pudo decodificar el archivo de texto")


def extraer_texto(nombre_archivo: str, contenido: bytes) -> str:
    nombre = (nombre_archivo or "").lower()
    if not nombre.endswith(EXTENSIONES):
        raise FormatoNoSoportado(
            "Formato no soportado: use PDF, DOCX o TXT (los .doc antiguos deben guardarse como .docx)"
        )
    if len(contenido) > TAMANO_MAXIMO:
        raise FormatoNoSoportado("El archivo supera el tamaño máximo de 5 MB")

    try:
        if nombre.endswith(".pdf"):
            texto = _pdf(contenido)
        elif nombre.endswith(".docx"):
            texto = _docx(contenido)
        else:
            texto = _txt(contenido)
    except (FormatoNoSoportado, ArchivoIlegible):
        raise
    except Exception as exc:
        raise ArchivoIlegible(f"No se pudo leer el archivo: {exc}") from exc

    texto = _limpiar(texto)
    if not texto:
        raise ArchivoIlegible(
            "El archivo no tiene texto extraíble (¿es un PDF escaneado? requiere OCR)"
        )
    return texto
